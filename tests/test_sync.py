"""Testes end-to-end do motor de sincronização."""

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_sync.sync import SyncConfig, sync


@pytest.fixture
def source_csv(tmp_path: Path) -> Path:
    path = tmp_path / "fonte.csv"
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Codigo", "Area", "Status"])
        writer.writerow(["TASK-001", "Financeiro", "Concluido"])
        writer.writerow(["TASK-002", "Juridico", "Em andamento"])
        writer.writerow(["TASK-999", "Comercial", "Pendente"])  # não existe no destino
    return path


@pytest.fixture
def target_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "destino.xlsx"
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "EquipeA"
    ws1.append(["Codigo", "Area", "Status"])
    ws1.append(["TASK-001", "Financeiro Antigo", "Aberto"])

    ws2 = wb.create_sheet("EquipeB")
    ws2.append(["Codigo", "Area", "Status"])
    ws2.append(["TASK-002", "Juridico", "Aberto"])  # área já igual, só status muda

    ws3 = wb.create_sheet("RESUMO")  # aba a ser ignorada
    ws3.append(["Codigo", "Area", "Status"])

    wb.save(path)
    return path


def test_sync_updates_only_changed_fields(source_csv, target_xlsx, tmp_path):
    log_path = tmp_path / "log.csv"
    config = SyncConfig(
        source_path=source_csv,
        target_path=target_xlsx,
        key_column="Codigo",
        update_columns=["Area", "Status"],
        header_row=1,
        ignored_sheets=["RESUMO"],
        log_path=log_path,
    )

    changes = sync(config)

    by_key = {c.key_value: c for c in changes}

    assert by_key["TASK-001"].status == "OK"
    assert by_key["TASK-001"].updated_columns == ["Area", "Status"]

    assert by_key["TASK-002"].status == "OK"
    assert by_key["TASK-002"].updated_columns == ["Status"]  # área já era igual

    assert by_key["TASK-999"].status == "NAO_ENCONTRADO"

    wb = load_workbook(target_xlsx)
    ws1 = wb["EquipeA"]
    assert ws1["B2"].value == "Financeiro"
    assert ws1["C2"].value == "Concluido"

    ws2 = wb["EquipeB"]
    assert ws2["B2"].value == "Juridico"
    assert ws2["C2"].value == "Em andamento"

    assert log_path.exists()


def test_sync_preserves_sheet_protection(source_csv, target_xlsx):
    wb = load_workbook(target_xlsx)
    wb["EquipeA"].protection.sheet = True
    wb["EquipeA"].protection.set_password("1234")
    wb.save(target_xlsx)

    config = SyncConfig(
        source_path=source_csv,
        target_path=target_xlsx,
        key_column="Codigo",
        update_columns=["Area", "Status"],
        header_row=1,
        ignored_sheets=["RESUMO"],
        sheet_password="1234",
    )
    sync(config)

    wb2 = load_workbook(target_xlsx)
    assert wb2["EquipeA"].protection.sheet is True
    assert wb2["EquipeA"]["B2"].value == "Financeiro"
