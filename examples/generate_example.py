"""
Gera um cenário de exemplo reproduzível para o excel-sync-tool.

Cenário (fictício): a "Nimbus Serviços" controla 160 processos internos em
uma planilha com uma aba por equipe. Todo dia o ERP exporta um CSV com o
status atualizado de cada processo. Este script cria:

- examples/controle.xlsx            planilha de controle (cabeçalho na linha 4,
                                    formatação, aba protegida e aba de resumo)
- examples/exportacao_sistema.csv   exportação "de hoje" do ERP

Uso:
    python examples/generate_example.py
    excel-sync --config examples/config.yaml
"""

from __future__ import annotations

import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).parent
SEED = 42
TEAMS = ["Financeiro", "Juridico", "Comercial", "RH"]
STATUSES = ["Aberto", "Em andamento", "Aguardando cliente", "Concluido"]
CLIENTS = [
    "Atlas Engenharia", "Bravo Logística", "Cedro Alimentos", "Delta Saúde",
    "Eixo Construtora", "Fênix Varejo", "Gama Têxtil", "Horizonte Agro",
]
OWNERS = ["Ana", "Bruno", "Carla", "Diego", "Elisa", "Fábio", "Gabi", "Hugo"]
HEADER = ["Codigo", "Cliente", "Responsavel", "Area", "Status", "Prazo"]
HEADER_ROW = 4
PASSWORD = "nimbus"


def build_processes(rng: random.Random, total: int = 160) -> list[dict[str, str]]:
    processes = []
    for n in range(1, total + 1):
        processes.append(
            {
                "Codigo": f"PRC-{n:04d}",
                "Cliente": rng.choice(CLIENTS),
                "Responsavel": rng.choice(OWNERS),
                "Area": TEAMS[(n - 1) % len(TEAMS)],
                "Status": rng.choice(STATUSES[:3]),
                "Prazo": f"2026-{rng.randint(10, 12):02d}-{rng.randint(1, 28):02d}",
            }
        )
    return processes


def build_daily_export(rng: random.Random, processes: list[dict[str, str]]) -> list[dict[str, str]]:
    """Simula o ERP: ~20% mudam de status, alguns trocam de área, 3 são novos."""
    export = []
    for p in processes:
        row = {"Codigo": p["Codigo"], "Area": p["Area"], "Status": p["Status"]}
        roll = rng.random()
        if roll < 0.20:
            row["Status"] = rng.choice([s for s in STATUSES if s != p["Status"]])
        elif roll < 0.25:
            row["Area"] = rng.choice([t for t in TEAMS if t != p["Area"]])
        export.append(row)
    for n in range(len(processes) + 1, len(processes) + 4):
        export.append({"Codigo": f"PRC-{n:04d}", "Area": rng.choice(TEAMS), "Status": "Aberto"})
    return export


def write_control_workbook(path: Path, processes: list[dict[str, str]]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "VISAO GERAL"
    summary["A1"] = "Nimbus Serviços — Controle de processos (dados fictícios)"
    summary["A1"].font = Font(bold=True, size=14)

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for team in TEAMS:
        ws = wb.create_sheet(team)
        ws["A1"] = f"Equipe {team}"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = "Atualizado automaticamente pelo excel-sync-tool"
        ws["A2"].font = Font(italic=True, color="666666")
        for col, name in enumerate(HEADER, start=1):
            cell = ws.cell(row=HEADER_ROW, column=col, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for p in (p for p in processes if p["Area"] == team):
            ws.append([p[h] for h in HEADER])
        for letter, width in zip("ABCDEF", (12, 20, 14, 14, 20, 12)):
            ws.column_dimensions[letter].width = width

    # Uma aba protegida para mostrar que a ferramenta desprotege e reprotege.
    wb["Juridico"].protection.set_password(PASSWORD)
    wb["Juridico"].protection.sheet = True

    wb.save(path)


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["Codigo", "Area", "Status"])
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    rng = random.Random(SEED)
    processes = build_processes(rng)
    write_control_workbook(HERE / "controle.xlsx", processes)
    write_csv(HERE / "exportacao_sistema.csv", build_daily_export(rng, processes))
    print(f"Gerado: {HERE / 'controle.xlsx'}")
    print(f"Gerado: {HERE / 'exportacao_sistema.csv'}")


if __name__ == "__main__":
    main()
